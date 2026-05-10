from pathlib import Path

from fastapi import APIRouter, BackgroundTasks, File, Form, UploadFile

from core.file_handler import output_path, save_upload
from core.job_queue import create_job, run_job

router = APIRouter()


def process_pdf_to_word(input_path: Path, job_id: str) -> Path:
    try:
        from pdf2docx import Converter
    except Exception as exc:
        raise NotImplementedError("pdf2docx is required for PDF to Word conversion") from exc

    out_path = output_path(job_id, ".docx")
    converter = Converter(str(input_path))
    try:
        converter.convert(str(out_path), start=0, end=None)
    finally:
        converter.close()
    return out_path


def process_pdf_to_txt(input_path: Path, job_id: str) -> Path:
    try:
        import fitz
    except Exception as exc:
        raise NotImplementedError("PyMuPDF is required for PDF text extraction") from exc

    out_path = output_path(job_id, ".txt")
    doc = fitz.open(str(input_path))
    text = "\n\n".join(page.get_text() for page in doc)
    out_path.write_text(text, encoding="utf-8")
    return out_path


def process_pdf_to_ppt(input_path: Path, job_id: str) -> Path:
    try:
        import fitz
        from pptx import Presentation
        from pptx.util import Inches
    except Exception as exc:
        raise NotImplementedError("PyMuPDF and python-pptx are required for PDF to PPT conversion") from exc

    out_path = output_path(job_id, ".pptx")
    doc = fitz.open(str(input_path))
    presentation = Presentation()
    blank = presentation.slide_layouts[6]
    for index, page in enumerate(doc):
        pix = page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5))
        image_path = output_path(f"{job_id}-{index}", ".png")
        pix.save(str(image_path))
        slide = presentation.slides.add_slide(blank)
        slide.shapes.add_picture(str(image_path), 0, 0, width=Inches(10))
    presentation.save(out_path)
    return out_path


def process_markup_to_pdf(input_path: Path, job_id: str, kind: str) -> Path:
    try:
        from weasyprint import HTML
        import markdown
    except Exception as exc:
        raise NotImplementedError("weasyprint and markdown are required for HTML/Markdown to PDF conversion") from exc

    html = input_path.read_text(encoding="utf-8", errors="ignore")
    if kind == "markdown":
        html = markdown.markdown(html)
    out_path = output_path(job_id, ".pdf")
    HTML(string=html).write_pdf(out_path)
    return out_path


async def _enqueue(background_tasks: BackgroundTasks, file: UploadFile, operation: str, processor) -> dict:
    job_id, input_path, original_filename = await save_upload(file)
    create_job("convert", operation, job_id, original_filename)
    background_tasks.add_task(run_job, job_id, lambda path: processor(path, job_id), input_path)
    return {"job_id": job_id, "status": "queued"}


@router.post("/pdf-to-word")
async def pdf_to_word(background_tasks: BackgroundTasks, file: UploadFile = File(...)) -> dict:
    return await _enqueue(background_tasks, file, "pdf-to-word", process_pdf_to_word)


@router.post("/pdf-to-txt")
async def pdf_to_txt(background_tasks: BackgroundTasks, file: UploadFile = File(...)) -> dict:
    return await _enqueue(background_tasks, file, "pdf-to-txt", process_pdf_to_txt)


@router.post("/pdf-to-ppt")
async def pdf_to_ppt(background_tasks: BackgroundTasks, file: UploadFile = File(...)) -> dict:
    return await _enqueue(background_tasks, file, "pdf-to-ppt", process_pdf_to_ppt)


@router.post("/html-to-pdf")
async def html_to_pdf(background_tasks: BackgroundTasks, file: UploadFile = File(...)) -> dict:
    return await _enqueue(background_tasks, file, "html-to-pdf", lambda path, job_id: process_markup_to_pdf(path, job_id, "html"))


@router.post("/markdown-to-pdf")
async def markdown_to_pdf(background_tasks: BackgroundTasks, file: UploadFile = File(...)) -> dict:
    return await _enqueue(background_tasks, file, "markdown-to-pdf", lambda path, job_id: process_markup_to_pdf(path, job_id, "markdown"))


@router.post("/pdf-to-excel")
async def pdf_to_excel(background_tasks: BackgroundTasks, file: UploadFile = File(...)) -> dict:
    return await _enqueue(background_tasks, file, "pdf-to-excel", lambda _path, _job_id: (_ for _ in ()).throw(NotImplementedError("camelot/pdfplumber table extraction is scaffolded but not enabled")))

@router.post("/pdf-to-html")
async def pdf_to_html(background_tasks: BackgroundTasks, file: UploadFile = File(...)) -> dict:
    return await _enqueue(background_tasks, file, "pdf-to-html", lambda _path, _job_id: (_ for _ in ()).throw(NotImplementedError("PDF to HTML is scaffolded but not enabled")))

@router.post("/pdf-to-odt")
async def pdf_to_odt(background_tasks: BackgroundTasks, file: UploadFile = File(...)) -> dict:
    return await _enqueue(background_tasks, file, "pdf-to-odt", lambda _path, _job_id: (_ for _ in ()).throw(NotImplementedError("PDF to ODT is scaffolded but not enabled")))

